import io
import sqlite3

from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.activity_feed import ActivityFeedStore
from app.approval_matrix import ALL_COMPANIES, ApprovalMatrixStore
from app.auth import AuthStore
from app.bulk_import import import_supplier_workbook
from app.companies import CompanyStore
from app.invoices import InvoiceStore
from app.main import create_app
from app.outlook_notifications import OutlookNotificationStore
from app.suppliers import SupplierStore
from app.supplier_terms import SupplierTermsStore


def client_for(tmp_path) -> TestClient:
    return TestClient(
        create_app(
            invoice_store=InvoiceStore(tmp_path / "invoices.db"),
            auth_store=AuthStore(tmp_path / "auth.db"),
            companies_store=CompanyStore(tmp_path / "config.db"),
            suppliers_store=SupplierStore(tmp_path / "config.db"),
            approval_matrix_store=ApprovalMatrixStore(tmp_path / "config.db"),
            supplier_terms_store=SupplierTermsStore(tmp_path / "config.db"),
            activity_feed=ActivityFeedStore(tmp_path / "activity_feed.db"),
            notification_store=OutlookNotificationStore(
                tmp_path / "outlook_notifications.db"
            ),
        )
    )


def login(client: TestClient, username: str, password: str) -> None:
    response = client.post(
        "/api/auth/login", json={"username": username, "password": password}
    )
    assert response.status_code == 200, response.text


def make_workbook(rows: list[list[object]]) -> io.BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "Company",
            "Supplier",
            "Supplier Account Number",
            "Default Payment Method",
            "Payment Terms",
            "Bank Account",
            "Approver",
            "Approver 2",
            "Approver Email",
            "Approver 2 Email",
        ]
    )
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def test_legacy_company_paths_are_normalized_to_current_structure(tmp_path) -> None:
    database_path = tmp_path / "config.db"
    with sqlite3.connect(database_path) as connection:
        connection.execute(
            """
            CREATE TABLE companies (
                name TEXT PRIMARY KEY,
                company_folder TEXT NOT NULL,
                po_matching_folder TEXT NOT NULL,
                aliases TEXT,
                vat_number TEXT,
                address TEXT
            )
            """
        )
        connection.execute(
            """
            INSERT INTO companies (
                name, company_folder, po_matching_folder, aliases
            ) VALUES (?, ?, ?, ?)
            """,
            (
                "Legacy Company",
                "Invoices/Legacy Company",
                "Invoices/Legacy Company/PO Matching",
                "",
            ),
        )

    company = CompanyStore(database_path).get("Legacy Company")

    assert company is not None
    assert company.sharepoint_root_folder == "Invoices/Legacy Company"
    assert company.company_folder == "Invoices/Legacy Company/Nominal Invoices"
    assert (
        company.po_matching_folder
        == "Invoices/Legacy Company/PO Invoices/PO Match"
    )


def test_import_supplier_workbook_creates_master_data_and_terms(tmp_path) -> None:
    company_store = CompanyStore(tmp_path / "config.db")
    supplier_store = SupplierStore(tmp_path / "config.db")
    approval_matrix_store = ApprovalMatrixStore(tmp_path / "config.db")
    supplier_terms_store = SupplierTermsStore(tmp_path / "config.db")
    workbook = make_workbook(
        [
            [
                "Brand New Co",
                "Brand New Supplier",
                "ACC-123",
                "BACS",
                "30 days net",
                "GBP Main Account",
                "Jordan Blake (Approver 1)",
                "Sam Ellis (Approver 2)",
                "jordan.blake@example.test",
                "sam.ellis@example.test",
            ],
            [
                "Brand New Co",
                "Unresolvable Supplier",
                "ACC-999",
                "BACS",
                "60 days net",
                "GBP Main Account",
                "Nobody Real",
                None,
            ],
        ]
    )

    summary = import_supplier_workbook(
        workbook,
        company_store=company_store,
        supplier_store=supplier_store,
        approval_matrix_store=approval_matrix_store,
        supplier_terms_store=supplier_terms_store,
    )

    assert summary.imported_count == 1
    assert summary.warning_count == 1
    assert summary.skipped_count == 0
    assert summary.rows[1].status == "warning"
    assert "Nobody Real" in summary.rows[1].reason

    assert company_store.get("Brand New Co") is not None
    imported_supplier = supplier_store.get("Brand New Supplier")
    assert imported_supplier is not None
    assert imported_supplier.default_company == "Brand New Co"

    matrix_entry = approval_matrix_store.find("Brand New Co", "Brand New Supplier")
    assert matrix_entry is not None
    assert matrix_entry.approver1.name == "Jordan Blake (Approver 1)"
    assert matrix_entry.approver2.name == "Sam Ellis (Approver 2)"

    terms = supplier_terms_store.get("Brand New Co", "Brand New Supplier")
    assert terms is not None
    assert terms.supplier_account_number == "ACC-123"
    assert terms.default_payment_method == "BACS"
    assert terms.payment_terms_notice == "30 days net"
    assert terms.bank_account == "GBP Main Account"

    # Supplier/payment and routing data remain useful before login accounts
    # and notification email addresses have been configured.
    assert supplier_store.get("Unresolvable Supplier") is not None
    assert (
        supplier_terms_store.get("Brand New Co", "Unresolvable Supplier")
        is not None
    )
    unresolved_route = approval_matrix_store.find_exact(
        "Brand New Co", "Unresolvable Supplier"
    )
    assert unresolved_route is not None
    assert unresolved_route.approver1.name == "Nobody Real"
    assert unresolved_route.approver1.email == ""


def test_admin_import_endpoint_requires_admin_role(tmp_path) -> None:
    client = client_for(tmp_path)
    login(client, "jordan.blake", "ChangeMe-App1!")

    workbook = make_workbook([])
    response = client.post(
        "/api/admin/import/supplier-master-data",
        files={
            "file": (
                "suppliers.xlsx",
                workbook.read(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 403


def test_admin_import_endpoint_imports_rows_end_to_end(tmp_path) -> None:
    client = client_for(tmp_path)
    login(client, "admin", "ChangeMe-Admin1!")

    workbook = make_workbook(
        [
            [
                "Acme Trading Ltd",
                "Supplier Ltd",
                "ACC-001",
                "Cheque",
                "14 days net",
                "GBP Acme Account",
                "Jordan Blake (Approver 1)",
                None,
                "jordan.blake@example.test",
                None,
            ]
        ]
    )
    workbook_bytes = workbook.read()
    duplicate_response = client.post(
        "/api/admin/import/supplier-master-data",
        data={"company": "Acme Trading Ltd"},
        files={
            "file": (
                "suppliers.xlsx",
                workbook_bytes,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert duplicate_response.status_code == 409
    duplicate_detail = duplicate_response.json()["detail"]
    assert duplicate_detail["duplicates"] == [
        {
            "existing_name": "Supplier Ltd",
            "spreadsheet_names": ["Supplier Ltd"],
            "row_numbers": [2],
        }
    ]
    assert (
        client.app.state.supplier_terms_store.get(
            "Acme Trading Ltd", "Supplier Ltd"
        )
        is None
    )

    response = client.post(
        "/api/admin/import/supplier-master-data",
        data={"company": "Acme Trading Ltd", "replace_existing": "true"},
        files={
            "file": (
                "suppliers.xlsx",
                workbook_bytes,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["imported"] == 1
    assert body["skipped"] == 0

    terms_response = client.get("/api/admin/supplier-terms")
    assert terms_response.status_code == 200
    terms = terms_response.json()
    assert any(
        t["supplier"] == "Supplier Ltd" and t["supplier_account_number"] == "ACC-001"
        for t in terms
    )
    payment_defaults = client.get(
        "/api/supplier-terms",
        params={"company": "Acme Trading Ltd", "supplier": "Supplier Ltd"},
    )
    assert payment_defaults.status_code == 200
    assert payment_defaults.json()["profiles"][0]["default_payment_method"] == "Cheque"


def test_import_confirmation_groups_duplicate_supplier_rows(tmp_path) -> None:
    client = client_for(tmp_path)
    login(client, "admin", "ChangeMe-Admin1!")
    workbook = make_workbook(
        [
            [
                None,
                "Supplier Ltd",
                "ACC-001",
                "BACS",
                "30 days",
                "GBP account",
                None,
                None,
            ],
            [
                None,
                "Supplier Limited",
                "ACC-002",
                "BACS",
                "30 days",
                "GBP account",
                None,
                None,
            ],
            [
                None,
                "New Supplier",
                "NEW-001",
                "BACS",
                "30 days",
                "GBP account",
                None,
                None,
            ],
        ]
    )
    workbook_bytes = workbook.read()

    response = client.post(
        "/api/admin/import/supplier-master-data",
        data={"company": "Acme Trading Ltd"},
        files={
            "file": (
                "suppliers.xlsx",
                workbook_bytes,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert response.status_code == 409
    assert response.json()["detail"]["duplicates"] == [
        {
            "existing_name": "Supplier Ltd",
            "spreadsheet_names": ["Supplier Ltd", "Supplier Limited"],
            "row_numbers": [2, 3],
        }
    ]
    assert client.app.state.suppliers_store.get("New Supplier") is None

    confirmed = client.post(
        "/api/admin/import/supplier-master-data",
        data={"company": "Acme Trading Ltd", "replace_existing": "true"},
        files={
            "file": (
                "suppliers.xlsx",
                workbook_bytes,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert confirmed.status_code == 200, confirmed.text
    imported_terms = client.app.state.supplier_terms_store.list_for_supplier(
        "Acme Trading Ltd", "Supplier Ltd"
    )
    assert {terms.supplier for terms in imported_terms} == {"Supplier Ltd"}


def test_selected_company_overrides_company_column_for_every_row(tmp_path) -> None:
    company_store = CompanyStore(tmp_path / "config.db")
    workbook = make_workbook(
        [
            [
                "Wrong Company",
                "Supplier One",
                "ONE-1",
                "BACS",
                "30 days",
                "Main account",
                None,
                None,
            ],
            [
                "Another Wrong Company",
                "Supplier Two",
                "TWO-1",
                "BANKLINE",
                "14 days",
                "Main account",
                None,
                None,
            ],
        ]
    )
    supplier_store = SupplierStore(tmp_path / "config.db")
    terms_store = SupplierTermsStore(tmp_path / "config.db")

    summary = import_supplier_workbook(
        workbook,
        company_store=company_store,
        supplier_store=supplier_store,
        approval_matrix_store=ApprovalMatrixStore(tmp_path / "config.db"),
        supplier_terms_store=terms_store,
        default_company="Acme Trading Ltd",
    )

    assert {row.company for row in summary.rows} == {"Acme Trading Ltd"}
    assert terms_store.get("Acme Trading Ltd", "Supplier One") is not None
    assert terms_store.get("Acme Trading Ltd", "Supplier Two") is not None
    assert supplier_store.get("Supplier One").default_company == "Acme Trading Ltd"
    assert supplier_store.get("Supplier Two").default_company == "Acme Trading Ltd"
    assert company_store.get("Wrong Company") is None
    assert company_store.get("Another Wrong Company") is None


def test_imports_attached_workbook_shape_for_selected_company(tmp_path) -> None:
    company_store = CompanyStore(tmp_path / "config.db")
    supplier_store = SupplierStore(tmp_path / "config.db")
    approval_matrix_store = ApprovalMatrixStore(tmp_path / "config.db")
    supplier_terms_store = SupplierTermsStore(tmp_path / "config.db")
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "index",
            "Supplier Account No",
            "Trading Partner Name",
            "DefaultPaymentMethod",
            "PaymentTerms",
            "Pay from bank account",
            "1st approval",
            "Final Signature",
        ]
    )
    sheet.append(
        [
            0,
            "ADOB0001",
            "Adobe Systems Software",
            "Direct Debit",
            "7 Days",
            "Nat West - Onecard",
            "Mark Kelly",
            "Graham Rogers",
        ]
    )
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    summary = import_supplier_workbook(
        buffer,
        company_store=company_store,
        supplier_store=supplier_store,
        approval_matrix_store=approval_matrix_store,
        supplier_terms_store=supplier_terms_store,
    )

    assert summary.warning_count == 1
    terms = supplier_terms_store.get(
        "Any Internal Company", "Adobe Systems Software"
    )
    assert terms is not None
    assert terms.company == ALL_COMPANIES
    assert terms.supplier_account_number == "ADOB0001"
    assert terms.default_payment_method == "Direct Debit"
    assert terms.payment_terms_notice == "7 Days"
    assert terms.bank_account == "Nat West - Onecard"
    assert "Mark Kelly" in summary.rows[0].reason
    route = approval_matrix_store.find_exact(
        ALL_COMPANIES, "Adobe Systems Software"
    )
    assert route is not None
    assert route.approver1.name == "Mark Kelly"
    assert route.approver1.email == ""
    assert route.approver2 is not None
    assert route.approver2.name == "Graham Rogers"
    assert route.approver2.email == ""


def test_selected_company_updates_existing_supplier_and_imports_named_approvers(
    tmp_path,
) -> None:
    company_store = CompanyStore(tmp_path / "config.db")
    company_store.create(
        name="GIFTED",
        sharepoint_root_folder="Invoices/GIFTED",
    )
    supplier_store = SupplierStore(tmp_path / "config.db")
    supplier_store.create(
        name="123RF GB Ltd",
        default_company="Acme Trading Ltd",
    )
    approval_matrix_store = ApprovalMatrixStore(tmp_path / "config.db")
    workbook = make_workbook(
        [
            [
                "Ignored workbook company",
                "123RF GB Ltd",
                "123R0001",
                "BACS",
                "30 days",
                "GBP account",
                "Jess Moon",
                "Julian Massie",
            ]
        ]
    )

    summary = import_supplier_workbook(
        workbook,
        company_store=company_store,
        supplier_store=supplier_store,
        approval_matrix_store=approval_matrix_store,
        supplier_terms_store=SupplierTermsStore(tmp_path / "config.db"),
        default_company="GIFTED",
    )

    assert supplier_store.get("123RF GB Ltd").default_company == "GIFTED"
    route = approval_matrix_store.find_exact("GIFTED", "123RF GB Ltd")
    assert route is not None
    assert route.approver1.name == "Jess Moon"
    assert route.approver1.email == ""
    assert route.approver2 is not None
    assert route.approver2.name == "Julian Massie"
    assert route.approver2.email == ""
    assert summary.warning_count == 1
    assert summary.rows[0].reason == (
        "Approval route imported. Add email addresses for: "
        "Jess Moon, Julian Massie."
    )

    approval_matrix_store.update(
        route.id,
        approver1_email="jess@example.test",
        approver2_email="julian@example.test",
    )
    workbook.seek(0)
    repeated = import_supplier_workbook(
        workbook,
        company_store=company_store,
        supplier_store=supplier_store,
        approval_matrix_store=approval_matrix_store,
        supplier_terms_store=SupplierTermsStore(tmp_path / "config.db"),
        default_company="GIFTED",
    )
    preserved = approval_matrix_store.find_exact("GIFTED", "123RF GB Ltd")
    assert preserved is not None
    assert preserved.approver1.email == "jess@example.test"
    assert preserved.approver2 is not None
    assert preserved.approver2.email == "julian@example.test"
    assert repeated.imported_count == 1


def test_company_specific_import_does_not_overwrite_global_route(tmp_path) -> None:
    company_store = CompanyStore(tmp_path / "config.db")
    supplier_store = SupplierStore(tmp_path / "config.db")
    approval_matrix_store = ApprovalMatrixStore(tmp_path / "config.db")
    approval_matrix_store.create(
        company=ALL_COMPANIES,
        supplier="Shared Supplier",
        approver1_name="Global Approver",
        approver1_email="global@example.test",
    )
    workbook = make_workbook(
        [
            [
                None,
                "Shared Supplier",
                "SHAR0001",
                "BACS",
                "30 days",
                "GBP account",
                "Jess Moon",
                None,
            ]
        ]
    )

    import_supplier_workbook(
        workbook,
        company_store=company_store,
        supplier_store=supplier_store,
        approval_matrix_store=approval_matrix_store,
        supplier_terms_store=SupplierTermsStore(tmp_path / "config.db"),
        default_company="Acme Trading Ltd",
    )

    global_route = approval_matrix_store.find_exact(
        ALL_COMPANIES, "Shared Supplier"
    )
    company_route = approval_matrix_store.find_exact(
        "Acme Trading Ltd", "Shared Supplier"
    )
    assert global_route is not None
    assert global_route.approver1.name == "Global Approver"
    assert company_route is not None
    assert company_route.approver1.name == "Jess Moon"


def test_global_approval_route_is_used_for_any_invoice_company(tmp_path) -> None:
    store = ApprovalMatrixStore(tmp_path / "config.db")
    store.create(
        company=ALL_COMPANIES,
        supplier="Adobe Systems Software",
        approver1_name="Mark Kelly",
        approver1_email="mark@example.test",
        approver2_name="Graham Rogers",
        approver2_email="graham@example.test",
    )

    route = store.find("Any Internal Company", "Adobe Systems Software")

    assert route is not None
    assert route.approver1.name == "Mark Kelly"
    assert route.approver2 is not None
    assert route.approver2.name == "Graham Rogers"


def test_workbook_without_company_imports_global_supplier_profile(tmp_path) -> None:
    workbook = Workbook()
    workbook.active.append(["Trading Partner Name"])
    workbook.active.append(["Adobe Systems Software"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    supplier_store = SupplierStore(tmp_path / "config.db")
    terms_store = SupplierTermsStore(tmp_path / "config.db")
    summary = import_supplier_workbook(
        buffer,
        company_store=CompanyStore(tmp_path / "config.db"),
        supplier_store=supplier_store,
        approval_matrix_store=ApprovalMatrixStore(tmp_path / "config.db"),
        supplier_terms_store=terms_store,
    )

    assert summary.warning_count == 1
    assert supplier_store.get("Adobe Systems Software") is not None
    profile = terms_store.list_for_supplier(
        "Any Internal Company", "Adobe Systems Software"
    )[0]
    assert profile.company == ALL_COMPANIES


def test_supplier_terms_store_migrates_legacy_company_supplier_key(tmp_path) -> None:
    database = tmp_path / "legacy.db"
    with sqlite3.connect(database) as connection:
        connection.execute(
            """
            CREATE TABLE supplier_terms (
                company TEXT NOT NULL,
                supplier TEXT NOT NULL,
                supplier_account_number TEXT,
                default_payment_method TEXT,
                payment_terms_notice TEXT,
                bank_account TEXT,
                PRIMARY KEY (company, supplier)
            )
            """
        )
        connection.execute(
            """
            INSERT INTO supplier_terms VALUES (?, ?, ?, ?, ?, ?)
            """,
            (
                "Acme Trading Ltd",
                "Supreme Freight Services Ltd",
                "SUPR0001",
                "BACS",
                "60 Days EOM",
                "Nat West - GBP",
            ),
        )

    store = SupplierTermsStore(database)
    store.upsert(
        company="Acme Trading Ltd",
        supplier="Supreme Freight Services Ltd",
        supplier_account_number="SUPR0002",
        default_payment_method="BACS",
        payment_terms_notice="60 Days EOM",
        bank_account="Nat West - USD",
    )

    profiles = store.list_for_supplier(
        "Acme Trading Ltd", "Supreme Freight Services Ltd"
    )
    assert [profile.supplier_account_number for profile in profiles] == [
        "SUPR0001",
        "SUPR0002",
    ]


def test_admin_can_edit_supplier_payment_settings(tmp_path) -> None:
    client = client_for(tmp_path)
    login(client, "admin", "ChangeMe-Admin1!")
    terms = client.app.state.supplier_terms_store.upsert(
        company="Acme Trading Ltd",
        supplier="Supplier Ltd",
        supplier_account_number="ACC-100",
        default_payment_method="BACS",
        payment_terms_notice="30 days",
    )

    response = client.put(
        f"/api/admin/supplier-terms/{terms.id}",
        json={
            "company": "Acme Trading Ltd",
            "supplier": "Supplier Ltd",
            "supplier_account_number": "ACC-200",
            "default_payment_method": "Bankline",
            "payment_terms_notice": None,
            "bank_account": "GBP current",
        },
    )

    assert response.status_code == 200, response.text
    assert response.json()["supplier_account_number"] == "ACC-200"
    assert response.json()["default_payment_method"] == "Bankline"
    assert response.json()["payment_terms_notice"] is None
    assert response.json()["bank_account"] == "GBP current"


def test_supplier_payment_update_rejects_duplicate_company_account(tmp_path) -> None:
    client = client_for(tmp_path)
    login(client, "admin", "ChangeMe-Admin1!")
    store = client.app.state.supplier_terms_store
    first = store.upsert(
        company="Acme Trading Ltd",
        supplier="Supplier One",
        supplier_account_number="ACC-100",
    )
    store.upsert(
        company="Acme Trading Ltd",
        supplier="Supplier Two",
        supplier_account_number="ACC-200",
    )

    response = client.put(
        f"/api/admin/supplier-terms/{first.id}",
        json={
            "company": "Acme Trading Ltd",
            "supplier": "Supplier One",
            "supplier_account_number": "ACC-200",
        },
    )

    assert response.status_code == 422
    assert "already exists" in response.json()["detail"]


def test_deleting_supplier_cascades_payment_profiles_and_routes(tmp_path) -> None:
    client = client_for(tmp_path)
    login(client, "admin", "ChangeMe-Admin1!")
    client.post("/api/admin/suppliers", json={"name": "Delete Me Ltd"})
    client.post(
        "/api/admin/approval-matrix",
        json={
            "company": "*",
            "supplier": "Delete Me Ltd",
            "approver1_name": "Jordan Blake",
            "approver1_email": "jordan@example.test",
        },
    )
    client.app.state.supplier_terms_store.upsert(
        company="*",
        supplier="Delete Me Ltd",
        supplier_account_number="DEL0001",
        default_payment_method="BACS",
    )

    response = client.delete("/api/admin/suppliers/Delete%20Me%20Ltd")

    assert response.status_code == 200
    assert client.app.state.approval_matrix_store.find(
        "Acme Trading Ltd", "Delete Me Ltd"
    ) is None
    assert client.app.state.supplier_terms_store.list_for_supplier(
        "Acme Trading Ltd", "Delete Me Ltd"
    ) == []


def test_deleting_internal_company_preserves_global_supplier_settings(tmp_path) -> None:
    client = client_for(tmp_path)
    login(client, "admin", "ChangeMe-Admin1!")
    client.post(
        "/api/admin/companies",
        json={
            "name": "Delete Company Ltd",
            "company_folder": "Invoices/Delete Company Ltd",
            "po_matching_folder": "Invoices/Delete Company Ltd/PO Matching",
        },
    )
    terms_store = client.app.state.supplier_terms_store
    terms_store.upsert(
        company="Delete Company Ltd",
        supplier="Supplier Ltd",
        supplier_account_number="SPECIFIC",
    )
    terms_store.upsert(
        company="*",
        supplier="Supplier Ltd",
        supplier_account_number="GLOBAL",
    )

    response = client.delete("/api/admin/companies/Delete%20Company%20Ltd")

    assert response.status_code == 200
    profiles = terms_store.list_for_supplier(
        "Delete Company Ltd", "Supplier Ltd"
    )
    assert [profile.supplier_account_number for profile in profiles] == ["GLOBAL"]
