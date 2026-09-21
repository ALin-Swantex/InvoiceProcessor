from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import BinaryIO

from openpyxl import load_workbook

from app.approval_matrix import ALL_COMPANIES, ApprovalMatrixStore
from app.companies import CompanyStore
from app.suppliers import SupplierProfile, SupplierStore
from app.supplier_terms import SupplierTermsStore


# ---------------------------------------------------------------------------
# Bulk import of company/supplier master data from a spreadsheet.
#
# This exists so an admin does not have to manually re-key every row of an
# existing Excel sheet (companies, supplier account numbers, default payment
# method, payment terms, bank account, and approver(s)) into the app one at
# a time. Accepted column headers are deliberately flexible (case/whitespace
# -insensitive, a few common synonyms) since real-world spreadsheets rarely
# match an exact schema.
#
# Rows may refer to approvers who do not have a local login yet. Their names
# are still imported into the matrix; a warning identifies routes that need
# an email address before outbound approval notifications are enabled.
# ---------------------------------------------------------------------------

_COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "company": ("company", "company name"),
    "supplier": (
        "supplier",
        "supplier name",
        "trading partner name",
    ),
    "supplier_account_number": (
        "supplier account number",
        "supplier account no",
        "account number",
        "supplier account",
    ),
    "default_payment_method": (
        "default payment method",
        "defaultpaymentmethod",
        "payment method",
    ),
    "payment_terms_notice": (
        "payment terms",
        "paymentterms",
        "payment terms notice",
        "terms",
    ),
    "bank_account": (
        "bank account",
        "bank account to pay from",
        "pay from bank account",
    ),
    "approver1": (
        "approver",
        "approver 1",
        "approver1",
        "1st approval",
        "approvers",
    ),
    "approver2": (
        "approver 2",
        "approver2",
        "second approver",
        "final signature",
    ),
    "approver1_email": (
        "approver email",
        "approver 1 email",
        "approver1 email",
        "1st approval email",
    ),
    "approver2_email": (
        "approver 2 email",
        "approver2 email",
        "second approver email",
        "final signature email",
    ),
}


@dataclass
class ImportRowResult:
    row_number: int
    company: str
    supplier: str
    status: str  # "imported", "warning", or "skipped"
    reason: str | None = None


@dataclass
class ImportSummary:
    rows: list[ImportRowResult] = field(default_factory=list)

    @property
    def imported_count(self) -> int:
        return sum(1 for row in self.rows if row.status == "imported")

    @property
    def skipped_count(self) -> int:
        return sum(1 for row in self.rows if row.status == "skipped")

    @property
    def warning_count(self) -> int:
        return sum(1 for row in self.rows if row.status == "warning")


@dataclass(frozen=True)
class ExistingSupplierImport:
    existing_name: str
    spreadsheet_names: tuple[str, ...]
    row_numbers: tuple[int, ...]


def _normalise_header(value: object) -> str:
    return str(value or "").strip().lower()


def _map_headers(header_row: tuple[object, ...]) -> dict[str, int]:
    """Map our canonical field names to the column index found in the sheet."""
    normalised = [_normalise_header(cell) for cell in header_row]
    mapping: dict[str, int] = {}
    for field_name, aliases in _COLUMN_ALIASES.items():
        for alias in aliases:
            if alias in normalised:
                mapping[field_name] = normalised.index(alias)
                break
    return mapping


def _cell(row: tuple[object, ...], mapping: dict[str, int], field_name: str) -> str | None:
    index = mapping.get(field_name)
    if index is None or index >= len(row):
        return None
    value = row[index]
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _split_names(raw: str) -> list[str]:
    """A single "Approvers" cell may list multiple names separated by a
    comma, semicolon, or the word "and"."""
    text = raw.replace(" and ", ",").replace(";", ",")
    return [part.strip() for part in text.split(",") if part.strip()]


def find_existing_supplier_imports(
    source: Path | BinaryIO,
    *,
    supplier_store: SupplierStore,
) -> list[ExistingSupplierImport]:
    """Return existing supplier companies referenced by a workbook.

    This is intentionally read-only so the API can request confirmation
    before any imported configuration is changed.
    """
    workbook = load_workbook(source, data_only=True, read_only=True)
    sheet = workbook.active
    rows_iter = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return []
    mapping = _map_headers(header_row)
    if "supplier" not in mapping:
        raise ValueError(
            "The spreadsheet must have a supplier column (for example "
            "'Supplier' or 'Trading Partner Name')."
        )

    supplier_lookup: dict[str, SupplierProfile] = {}
    for profile in supplier_store.list():
        supplier_lookup[profile.name.strip().casefold()] = profile
        for alias in profile.aliases:
            supplier_lookup[alias.strip().casefold()] = profile

    duplicates: dict[str, dict[str, object]] = {}
    for row_number, row in enumerate(rows_iter, start=2):
        supplier = _cell(row, mapping, "supplier")
        if not supplier:
            continue
        existing = supplier_lookup.get(supplier.strip().casefold())
        if existing is None:
            continue
        key = existing.name.strip().casefold()
        duplicate = duplicates.setdefault(
            key,
            {
                "existing_name": existing.name,
                "spreadsheet_names": [],
                "row_numbers": [],
            },
        )
        spreadsheet_names = duplicate["spreadsheet_names"]
        row_numbers = duplicate["row_numbers"]
        assert isinstance(spreadsheet_names, list)
        assert isinstance(row_numbers, list)
        if supplier not in spreadsheet_names:
            spreadsheet_names.append(supplier)
        row_numbers.append(row_number)

    return [
        ExistingSupplierImport(
            existing_name=str(duplicate["existing_name"]),
            spreadsheet_names=tuple(duplicate["spreadsheet_names"]),
            row_numbers=tuple(duplicate["row_numbers"]),
        )
        for duplicate in duplicates.values()
    ]


def import_supplier_workbook(
    source: Path | BinaryIO,
    *,
    company_store: CompanyStore,
    supplier_store: SupplierStore,
    approval_matrix_store: ApprovalMatrixStore,
    supplier_terms_store: SupplierTermsStore,
    default_company: str | None = None,
) -> ImportSummary:
    workbook = load_workbook(source, data_only=True, read_only=True)
    sheet = workbook.active

    rows_iter = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return ImportSummary()
    mapping = _map_headers(header_row)

    if "supplier" not in mapping:
        raise ValueError(
            "The spreadsheet must have a supplier column (for example "
            "'Supplier' or 'Trading Partner Name')."
        )
    selected_company = (default_company or "").strip()
    if selected_company and company_store.get(selected_company) is None:
        raise ValueError(f"Selected company '{selected_company}' was not found.")
    bulk_writer = getattr(supplier_store, "bulk_import_master_data", None)
    use_bulk_writer = callable(bulk_writer) and bool(selected_company)
    supplier_lookup: dict[str, SupplierProfile] = {}
    for profile in supplier_store.list():
        supplier_lookup[profile.name.strip().casefold()] = profile
        for alias in profile.aliases:
            supplier_lookup[alias.strip().casefold()] = profile
    approval_lookup = {
        (entry.company.strip().casefold(), entry.supplier.strip().casefold()): entry
        for entry in approval_matrix_store.list()
    }
    bulk_suppliers: dict[str, tuple[str, str]] = {}
    bulk_approvals: dict[
        tuple[str, str],
        tuple[str, str, str, str, str | None, str | None],
    ] = {}
    bulk_terms: dict[
        tuple[str, str],
        tuple[
            str,
            str,
            str | None,
            str | None,
            str | None,
            str | None,
        ],
    ] = {}

    summary = ImportSummary()
    for row_number, row in enumerate(rows_iter, start=2):
        if row is None or all(cell is None for cell in row):
            continue

        # An admin-selected company scopes the entire workbook and takes
        # precedence over any Company column accidentally left in the file.
        company = selected_company or _cell(row, mapping, "company") or ALL_COMPANIES
        supplier = _cell(row, mapping, "supplier")
        if not company or not supplier:
            summary.rows.append(
                ImportRowResult(
                    row_number=row_number,
                    company=company or "",
                    supplier=supplier or "",
                    status="skipped",
                    reason="Missing company or supplier name.",
                )
            )
            continue
        supplier_profile = supplier_lookup.get(supplier.strip().casefold())
        if supplier_profile is not None:
            supplier = supplier_profile.name

        approver1_raw = _cell(row, mapping, "approver1")
        approver2_raw = _cell(row, mapping, "approver2")
        approver_names = _split_names(approver1_raw) if approver1_raw else []
        if approver2_raw:
            approver_names.append(approver2_raw)

        supplied_emails = [
            _cell(row, mapping, "approver1_email"),
            _cell(row, mapping, "approver2_email"),
        ]
        existing_entry = approval_lookup.get(
            (company.strip().casefold(), supplier.strip().casefold())
        )
        existing_emails = (
            {
                approver.name.strip().casefold(): approver.email
                for approver in (
                    existing_entry.approver1,
                    existing_entry.approver2,
                )
                if approver is not None and approver.email
            }
            if existing_entry is not None
            else {}
        )
        resolved: list[tuple[str, str]] = []
        missing_emails: list[str] = []
        for index, name in enumerate(approver_names[:2]):
            display_name = name
            email = (
                supplied_emails[index]
                or existing_emails.get(display_name.strip().casefold())
            )
            resolved.append((display_name, email or ""))
            if not email:
                missing_emails.append(display_name)
        route_warning: str | None = None
        if not approver1_raw:
            route_warning = "No first approver was supplied."
        elif missing_emails:
            route_warning = (
                "Approval route imported. Add email addresses for: "
                f"{', '.join(missing_emails)}."
            )

        # Ensure the company and supplier master records exist so the
        # approval matrix / supplier terms rows can reference them; leave
        # folder paths for the admin to adjust afterwards if this created a
        # brand-new company.
        if (
            not use_bulk_writer
            and company != ALL_COMPANIES
            and company_store.get(company) is None
        ):
            company_store.create(
                name=company,
                sharepoint_root_folder=f"Invoices/{company}",
            )
        supplier_default_company = company if company != ALL_COMPANIES else None
        if use_bulk_writer:
            bulk_suppliers[supplier.casefold()] = (
                supplier,
                supplier_default_company or "",
            )
        elif supplier_profile is None:
            supplier_store.create(
                name=supplier,
                default_company=supplier_default_company,
            )
        elif supplier_default_company:
            supplier_store.update(
                supplier_profile.name,
                default_company=supplier_default_company,
            )

        if resolved:
            approver1_name, approver1_email = resolved[0]
            approver2_name, approver2_email = resolved[1] if len(resolved) > 1 else (None, None)
            if use_bulk_writer:
                bulk_approvals[
                    (company.casefold(), supplier.casefold())
                ] = (
                    company,
                    supplier,
                    approver1_name,
                    approver1_email,
                    approver2_name,
                    approver2_email,
                )
            elif existing_entry is None:
                approval_matrix_store.create(
                    company=company,
                    supplier=supplier,
                    approver1_name=approver1_name,
                    approver1_email=approver1_email,
                    approver2_name=approver2_name,
                    approver2_email=approver2_email,
                )
            else:
                approval_matrix_store.update(
                    existing_entry.id,
                    approver1_name=approver1_name,
                    approver1_email=approver1_email,
                    approver2_name=approver2_name,
                    approver2_email=approver2_email,
                )

        account_number = _cell(row, mapping, "supplier_account_number")
        term_values = (
            company,
            supplier,
            account_number,
            _cell(row, mapping, "default_payment_method"),
            _cell(row, mapping, "payment_terms_notice"),
            _cell(row, mapping, "bank_account"),
        )
        if use_bulk_writer:
            bulk_terms[
                (
                    company.casefold(),
                    account_number.casefold()
                    if account_number
                    else supplier.casefold(),
                )
            ] = term_values
        else:
            supplier_terms_store.upsert(
                company=company,
                supplier=supplier,
                supplier_account_number=account_number,
                default_payment_method=term_values[3],
                payment_terms_notice=term_values[4],
                bank_account=term_values[5],
            )

        summary.rows.append(
            ImportRowResult(
                row_number=row_number,
                company=company,
                supplier=supplier,
                status="warning" if route_warning else "imported",
                reason=route_warning,
            )
        )

    if use_bulk_writer:
        bulk_writer(
            suppliers=list(bulk_suppliers.values()),
            approvals=list(bulk_approvals.values()),
            terms=list(bulk_terms.values()),
        )
    return summary
